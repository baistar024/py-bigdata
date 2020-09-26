import re, os
import openpyxl, docx
from win32com import client as wc
#---------------------------------------------------------------------
#将指定文件夹下的文件检索出来，并返回文件列表
def getfiles(subpath = "", extname = ".doc" ):
    path = os.path.join(os.getcwd(), subpath)
    files = []
    for file in os.listdir(path):
        if file.endswith(extname):
            files.append(os.path.join(path, file))
    return files
def removedoc(files):
    for file in files:
        os.remove(file)
#-------------------------------------------------------------------
#将指定的doc文件转换为docx文件
def doc2docx(files=[]):
    word = wc.Dispatch("Word.Application")
    for file in files:
        doc = word.Documents.Open(file)
        doc.SaveAs("{}x".format(file), 12)
        doc.Close()
    word.Quit()
#-----------------------------------------------------------------
# 从字符串中提取检查日期
def getdate(text = ""):
    datepattern = r'2020.\s*[0|1]?\d.\s*[0|1|2|3]?\d'
    match = re.search(datepattern, text, re.I).group()
    return match
#--------------------------------------------------------------------
# 从字符串中提取楼号信息
def getdorm(text = ""):
    dormpattern = r'(1[4|5|7|8]号)|(楼号：1[4|5|7|8])|(1[4|5|6|7|8]号楼)'
    findtext =re.search(dormpattern, text, re.I).group()
    match = re.search(r'\d\d', findtext).group()
    return match
#----------------------------------------------------------------
# 从台头信息中找到检查日期和楼号
def gettabhead(file):
    # dorms = [14, 15, 17, 18]
    heads = []
    texts = getdocxtext(file)
    date = getdate(texts[0])
    for text in texts:
        dormno = getdorm(text)
        heads.append((date, dormno))
        print(date, dormno)
    return heads
# ------------------------------------------------------------------
# function;getdocxtext(docfile)
# docfile : word文档名称
# return :返回表关前的台头信息
def getdocxtext(docfile):
    doc = docx.Document(docfile)
    texts = []
    for para in doc.paragraphs:
        txt = para.text
        if ("日期" in txt or ("日" in txt and "月" in txt)) and ("楼号" in txt or "号楼" in txt or "号" in txt):
            texts.append(para.text)
    return texts

def gettabinfo(files = [], department = "艺术"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["检查日期","楼号","寝室号","系别","扣分原因","分数"])
    for file in files:
        print(file)
        heads = gettabhead(file)
        doc = docx.Document(file)
        tabs = doc.tables
        i = 0
        for tab in tabs:
            for row in tab.rows:
                row_content = []
                row_content.append(heads[i][0])
                row_content.append(heads[i][1])
                if row.cells[1].text == department:
                    for cell in row.cells:
                        row_content.append(cell.text)
                    row_content = dataclearn(row_content)
                    ws.append(row_content)


            i += 1
    wb.save("Score-2020-2021A.xlsx")
    wb.close()

def dataclearn(items):
    items[0] = items[0].replace('.','/')
    if items[5]:
        items[5] = re.search( r'\d?\d', items[5]).group()
    return items

# -------------------------------------------------------------------
#主程序　
#1.找出所有指定目录下的doc文件,将所有的doc文件转换成docx文件
# 2.处理每一个docx文件并写入到excel文件中
# 2.1在当前的docx文件中找出日期和楼号
# 2.2列出所有的表格并在表格中找出有用的数据
# 2.3找到有用的数据后写入到excel文件中

def main():
    subpath = 'record'
    extdoc = '.doc'
    extdocx = '.docx'
    files = getfiles(subpath= subpath)
    if(len(files))>0:
        doc2docx(files)
        removedoc(files)
    files =getfiles(subpath= subpath, extname= extdocx)
    if len(files) == 0:
        return
    else:
      gettabinfo(files = files)
    print("data process is finish!")
main()



